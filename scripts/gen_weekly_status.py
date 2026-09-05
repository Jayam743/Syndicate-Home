#!/usr/bin/env python3
"""Deterministic renderer for the weekly status .docx report.

Content comes in as a JSON payload (produced by Ledger via the /status skill);
this script only renders it. It does NOT invent or synthesize content.
"""
import argparse
import json
import os
import sys
import tempfile

from docx import Document
from docx.shared import RGBColor

BLACK = RGBColor(0x00, 0x00, 0x00)

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFIG_PATH = os.path.join(REPO_ROOT, "config", "weekly-status.json")

REQUIRED_KEYS = (
    "week_range_label",
    "filename",
    "accomplishments",
    "blockers",
    "track",
    "next_obj",
)


def add_bold_line(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = BLACK
    return p


def add_italic_line(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = BLACK
    return p


def add_section_label(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = BLACK
    return p


def add_theme_heading(doc, text):
    """○-level bold theme lead-in, e.g. 'Cotterpin engagement lifecycle (agent-mantle):'"""
    p = doc.add_paragraph(style="List Bullet 2")
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = BLACK
    return p


def add_detail_bullet(doc, text):
    """▪-level crisp point under a theme."""
    p = doc.add_paragraph(style="List Bullet 3")
    r = p.add_run(text)
    r.font.color.rgb = BLACK
    return p


def add_plain_bullet(doc, text, level="List Bullet 2"):
    p = doc.add_paragraph(style=level)
    r = p.add_run(text)
    r.font.color.rgb = BLACK
    return p


def build_report(out_dir, filename, week_range_label, accomplishments, blockers, track, next_obj):
    """
    accomplishments: list of (theme_heading, [detail1, detail2, ...])
    """
    doc = Document()

    add_bold_line(doc, "Patel, Jayam")
    add_italic_line(doc, week_range_label)

    add_section_label(doc, "Accomplishments:")
    for theme, details in accomplishments:
        add_theme_heading(doc, theme)
        for d in details:
            add_detail_bullet(doc, d)

    add_section_label(doc, "Blockers, and whether or not there is a mitigation:")
    add_plain_bullet(doc, blockers)

    add_section_label(doc, "Are you on track, behind, or ahead? By how much?")
    add_plain_bullet(doc, track)

    add_section_label(doc, "Next week's objectives:")
    add_plain_bullet(doc, next_obj)

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, filename)
    doc.save(path)
    return path


def resolve_out_dir():
    """Config out_dir, overridable via SYNDICATE_STATUS_OUT_DIR, with ~ expansion."""
    override = os.environ.get("SYNDICATE_STATUS_OUT_DIR")
    if override:
        return os.path.expanduser(override)
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        cfg = json.load(f)
    out_dir = cfg.get("out_dir")
    if not out_dir:
        raise SystemExit(f"config {CONFIG_PATH} missing non-empty 'out_dir'")
    return os.path.expanduser(out_dir)


def load_payload(payload_path):
    with open(payload_path, "r", encoding="utf-8") as f:
        payload = json.load(f)

    if not isinstance(payload, dict):
        raise SystemExit("payload must be a JSON object")

    keys = set(payload.keys())
    missing = [k for k in REQUIRED_KEYS if k not in keys]
    extra = [k for k in keys if k not in REQUIRED_KEYS]
    empty = [k for k in REQUIRED_KEYS if k in keys and not payload[k]]
    problems = []
    if missing:
        problems.append(f"missing keys: {missing}")
    if extra:
        problems.append(f"unexpected keys: {extra}")
    if empty:
        problems.append(f"empty keys: {empty}")
    if problems:
        raise SystemExit(
            "invalid payload "
            + payload_path
            + " — "
            + "; ".join(problems)
            + f". Required (all non-empty): {list(REQUIRED_KEYS)}"
        )

    accomplishments = []
    for i, item in enumerate(payload["accomplishments"]):
        if not isinstance(item, dict) or "theme" not in item or "details" not in item:
            raise SystemExit(
                f"accomplishments[{i}] must be an object with 'theme' and 'details'"
            )
        if not isinstance(item["theme"], str) or not isinstance(item["details"], list):
            raise SystemExit(
                f"accomplishments[{i}] 'theme' must be a string and 'details' a list"
            )
        if not item["theme"] or not item["details"]:
            raise SystemExit(f"accomplishments[{i}] has empty 'theme' or 'details'")
        accomplishments.append((item["theme"], list(item["details"])))
    payload["_accomplishments_tuples"] = accomplishments
    return payload


def render(payload_path, out_dir=None):
    payload = load_payload(payload_path)
    if out_dir is None:
        out_dir = resolve_out_dir()
    path = build_report(
        out_dir,
        payload["filename"],
        payload["week_range_label"],
        payload["_accomplishments_tuples"],
        payload["blockers"],
        payload["track"],
        payload["next_obj"],
    )
    print(path)
    return path


def render_text(payload_path):
    """Terminal text preview of the payload content (no docx written)."""
    payload = load_payload(payload_path)
    lines = []
    lines.append("Patel, Jayam")
    lines.append(payload["week_range_label"])
    lines.append("")
    lines.append("Accomplishments:")
    for theme, details in payload["_accomplishments_tuples"]:
        lines.append(f"  - {theme}")
        for d in details:
            lines.append(f"      - {d}")
    lines.append("")
    lines.append("Blockers, and whether or not there is a mitigation:")
    lines.append(f"  - {payload['blockers']}")
    lines.append("")
    lines.append("Are you on track, behind, or ahead? By how much?")
    lines.append(f"  - {payload['track']}")
    lines.append("")
    lines.append("Next week's objectives:")
    lines.append(f"  - {payload['next_obj']}")
    print("\n".join(lines))


SAMPLE_PAYLOAD = {
    "week_range_label": "Week of Thu Aug 27 – Wed Sep 2, 2026",
    "filename": "Weekly Status - Patel, Jayam - Week ending 2026-09-02.docx",
    "accomplishments": [
        {
            "theme": "Auth service hardening (example-api):",
            "details": [
                "Built token refresh with rotation and a leak-proof session store (#238)",
                "Shipped a scriptable admin CLI (#240) and approval checks (#241)",
            ],
        },
        {
            "theme": "Config-as-code (example-infra):",
            "details": [
                "Shipped the bootstrap playbook — “config-as-code” for the substrate",
                "Migrated the app DB from SQLite to Postgres with a retention cap",
            ],
        },
    ],
    "blockers": "None active — the within-week correctness bug (#269) was caught and fixed the same week.",
    "track": "Ahead. 55 MRs merged across three repos this week — up from 34 last week.",
    "next_obj": "Continue hardening the engagement lifecycle and build out Salvo’s config-as-code surface.",
}


def smoke():
    tmp_dir = tempfile.mkdtemp(prefix="status-smoke-")
    payload_path = os.path.join(tmp_dir, "sample-payload.json")
    with open(payload_path, "w", encoding="utf-8") as f:
        json.dump(SAMPLE_PAYLOAD, f, ensure_ascii=False)

    path = render(payload_path, out_dir=tmp_dir)

    size = os.path.getsize(path)
    if size == 0:
        print(f"SMOKE FAIL: rendered docx is zero bytes: {path}", file=sys.stderr)
        return 1

    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    if "Patel, Jayam" not in texts:
        print("SMOKE FAIL: rendered docx missing expected content", file=sys.stderr)
        return 1

    print(f"OK ({size} bytes, {len(texts)} paragraphs): {path}")
    return 0


def main():
    parser = argparse.ArgumentParser(description="Deterministic weekly status .docx renderer.")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--payload", metavar="FILE", help="render a JSON payload to a .docx")
    group.add_argument("--text", metavar="FILE", help="print a terminal text preview of a payload")
    group.add_argument("--smoke", action="store_true", help="render a built-in sample and self-check")
    args = parser.parse_args()

    if args.smoke:
        sys.exit(smoke())
    if args.text:
        render_text(args.text)
        sys.exit(0)
    if args.payload:
        render(args.payload)
        sys.exit(0)


if __name__ == "__main__":
    main()
